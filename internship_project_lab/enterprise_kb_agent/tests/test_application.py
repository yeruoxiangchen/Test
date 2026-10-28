import io
import json
from pathlib import Path
import pytest
from kb.agent import Harness
from kb.ingestion import ingest, register_document
from kb.retrieval import Retriever


def upload(system, content="出差住宿标准每晚 600 元。报销申请需在出差结束后 15 天提交。", name="报销.md", private=False):
    cfg, db, client, public, restricted, *_ = system
    response = client.post("/api/documents", data={"base_id":restricted if private else public}, files={"file":(name, content.encode() if isinstance(content,str) else content)})
    assert response.status_code == 202, response.text
    doc = response.json()
    ingest(db, cfg, doc["id"])
    return doc


def login(client, username="member"):
    return client.post("/api/login", json={"username":username, "password":f"test-{username}-123"})


def test_upload_retrieve_answer_sources_and_feedback(system):
    cfg, db, client, base, *_ = system
    doc = upload(system)
    assert client.get("/api/documents").json()[0]["status"] == "ready"
    login(client)
    response = client.post("/api/chat", json={"question":"出差住宿报销标准是多少？", "base_ids":[base]})
    assert response.status_code == 202
    run_id = response.json()["run_id"]
    Harness(db,cfg).execute(run_id)
    run = client.get("/api/runs/"+run_id).json()
    assert run["status"] == "answered" and "600" in run["answer"]
    assert run["citations"][0]["document_id"] == doc["id"]
    assert "未调用大模型" in run["answer"]
    assert client.post(f"/api/runs/{run_id}/feedback", json={"rating":1,"comment":"找到原文了"}).status_code == 200
    stream = client.get(f"/api/runs/{run_id}/events").text
    assert "event: done" in stream and "read_chunk" in stream
    login(client,"admin")
    assert client.get("/api/feedback").json()[0]["comment"] == "找到原文了"


def test_member_cannot_read_admin_documents_or_mutate(system):
    cfg, db, client, base, private, admin, member = system
    doc = upload(system,"秘密奖金金额为 987654。",private=True)
    chunk_id = db.one("SELECT id FROM chunks WHERE document_id=?",(doc["id"],))["id"]
    login(client)
    assert all(b["id"] != private for b in client.get("/api/bases").json())
    assert client.get("/api/documents/"+doc["id"]).status_code == 403
    assert client.get("/api/documents/"+doc["id"]+"/download").status_code == 403
    assert client.get("/api/search",params={"q":"奖金"}).json()["results"] == []
    assert client.post("/api/chat",json={"question":"奖金", "base_ids":[private]}).status_code == 403
    assert client.post("/api/bases",json={"name":"越权"}).status_code == 403
    assert client.delete("/api/documents/"+doc["id"]).status_code == 403
    assert client.get("/api/feedback").status_code == 403
    with pytest.raises(PermissionError):
        Retriever(db,cfg).chunk(member,chunk_id,[base])


def test_duplicate_delete_removes_fts_and_vectors(system):
    cfg, db, client, base, *_ = system
    first = upload(system)
    response = client.post("/api/documents", data={"base_id":base},files={"file":("同内容.md","出差住宿标准每晚 600 元。报销申请需在出差结束后 15 天提交。".encode())})
    assert response.json()["duplicate"] and response.json()["id"] == first["id"]
    chunk = db.one("SELECT id FROM chunks WHERE document_id=?",(first["id"],))["id"]
    db.execute("INSERT INTO vectors VALUES(?,?,?)",(chunk,"profile","[1,0]"))
    assert client.delete("/api/documents/"+first["id"]).status_code == 200
    assert db.one("SELECT count(*) AS n FROM chunks_fts")["n"] == 0
    assert db.one("SELECT count(*) AS n FROM vectors")["n"] == 0
    assert not (cfg.root/"uploads"/first["id"]).exists()
    assert client.get("/api/search",params={"q":"住宿"}).json()["results"] == []


def test_reindex_replaces_without_duplicate_chunks(system):
    cfg, db, client, *_ = system
    doc=upload(system)
    count=db.one("SELECT count(*) AS n FROM chunks")["n"]
    assert client.post("/api/documents/"+doc["id"]+"/reindex").status_code==200
    ingest(db,cfg,doc["id"])
    assert db.one("SELECT count(*) AS n FROM chunks")["n"] == count
    assert db.one("SELECT count(*) AS n FROM chunks_fts")["n"] == count


def test_docx_tables_and_pdf_pages(system):
    from docx import Document
    doc=Document();doc.add_paragraph("设备保修政策")
    table=doc.add_table(rows=1,cols=2);table.cell(0,0).text="保修期限";table.cell(0,1).text="12 个月"
    buffer=io.BytesIO();doc.save(buffer)
    result=upload(system,buffer.getvalue(),"设备.docx")
    cfg,db,client,*_=system
    assert "12 个月" in client.get("/api/documents/"+result["id"]).json()["chunks"][0]["text"]
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer=PdfWriter();page=writer.add_blank_page(width=600,height=800)
    font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
    page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
    stream=DecodedStreamObject();stream.set_data(b'BT /F1 12 Tf 50 700 Td (Travel reimbursement limit 600.) Tj ET')
    page[NameObject('/Contents')]=writer._add_object(stream)
    buffer=io.BytesIO();writer.write(buffer)
    pdf=upload(system,buffer.getvalue(),"travel.pdf")
    parts=client.get("/api/documents/"+pdf["id"]).json()["chunks"]
    assert parts[0]["page"] == 1 and "600" in parts[0]["text"]


def test_bad_or_empty_file_has_explicit_error(system):
    cfg,db,client,base,*_=system
    bad=upload(system,b"not a valid PDF","bad.pdf")
    assert db.one("SELECT status FROM documents WHERE id=?",(bad["id"],))["status"] == "error"
    empty=client.post("/api/documents",data={"base_id":base},files={"file":("empty.txt",b"")})
    assert empty.status_code == 400
    executable=client.post("/api/documents",data={"base_id":base},files={"file":("run.py",b"print(1)")})
    assert executable.status_code == 400
    traversal=upload(system,"正常文档内容","../../escape.md")
    assert traversal["filename"] == "escape.md" and not (cfg.root/"escape.md").exists()


def test_session_isolation_and_login_logout(system):
    cfg,db,client,base,*_=system
    upload(system)
    login(client)
    created=client.post("/api/chat",json={"question":"住宿", "base_ids":[base]}).json()
    login(client,"second")
    assert client.get("/api/runs/"+created["run_id"]).status_code == 404
    assert client.get("/api/conversations/"+created["conversation_id"]).status_code == 404
    assert client.post("/api/chat",json={"question":"接着回答", "conversation_id":created["conversation_id"]}).status_code == 404
    assert client.post("/api/logout").status_code==200
    assert client.get("/api/me").status_code==401
    assert client.post("/api/login",json={"username":"admin","password":"incorrect"}).status_code==401


def test_csrf_and_upload_limit(system):
    cfg,db,client,base,*_=system
    assert client.post("/api/bases",json={"name":"bad"},headers={"Origin":"https://other.invalid"}).status_code==403
    assert client.post("/api/bases",json={"name":"bad"},headers={"X-KB-Client":""}).status_code==403
    assert client.post("/api/documents",content=b"x",headers={"Content-Length":str(cfg.max_upload_bytes+1024*1024+1)}).status_code==413


def test_no_results_and_cancel_and_one_run_per_conversation(system):
    cfg,db,client,base,*_=system
    upload(system)
    created=client.post("/api/chat",json={"question":"火星地下矿产的具体分布", "base_ids":[base]}).json()
    assert client.post("/api/chat",json={"question":"第二问", "conversation_id":created["conversation_id"]}).status_code==409
    Harness(db,cfg).execute(created["run_id"])
    assert client.get("/api/runs/"+created["run_id"]).json()["status"]=="insufficient"
    queued=client.post("/api/chat",json={"question":"住宿"}).json()
    assert client.post("/api/runs/"+queued["run_id"]+"/cancel").status_code==200
    Harness(db,cfg).execute(queued["run_id"])
    assert client.get("/api/runs/"+queued["run_id"]).json()["status"]=="cancelled"


def test_embedding_profile_and_access_filters(system,monkeypatch):
    cfg,db,client,base,private,admin,member=system
    public_doc=upload(system)
    private_doc=upload(system,"机密奖金金额 987654",private=True)
    public_chunk=db.one("SELECT id FROM chunks WHERE document_id=?",(public_doc["id"],))["id"]
    private_chunk=db.one("SELECT id FROM chunks WHERE document_id=?",(private_doc["id"],))["id"]
    cfg.embedding_url="http://embedding.invalid/v1";cfg.embedding_model="example"
    db.execute("INSERT INTO vectors VALUES(?,?,?)",(public_chunk,"wrong-profile","[1,0]"))
    db.execute("INSERT INTO vectors VALUES(?,?,?)",(private_chunk,cfg.embedding_profile,"[1,0]"))
    monkeypatch.setattr("kb.retrieval.embedding_vectors",lambda s,t:[[1,0]])
    results,warnings=Retriever(db,cfg).search(member,"差旅花费",[base])
    assert results==[] and warnings
    db.execute("INSERT INTO vectors VALUES(?,?,?)",(public_chunk,cfg.embedding_profile,"[1,0]"))
    results,_=Retriever(db,cfg).search(member,"差旅花费",[base])
    assert len(results)==1 and results[0]["document_id"]==public_doc["id"]
